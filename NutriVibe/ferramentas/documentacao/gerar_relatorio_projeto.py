"""Redige o relatório do projeto a partir do modelo, sem conservar listas vazias."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import sys,json
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
MODEL=Path(sys.argv[1]);OUT=ROOT/'documentacoes/relatorio-projeto';OUT.mkdir(exist_ok=True)
doc=Document(MODEL);original=list(doc.paragraphs)
headings=[deepcopy(original[i]._p)for i in [18,32,47,64,74,93,113,125,139]]
cover={0:'PROJETO: NutriVibe',4:'Turma: 101',7:'Nome do Grupo: NutriVibe',8:'Paulo Capreti — RA: 26028858',9:'Davi Pioli — RA: 26005199',10:'Victor Almeida — RA: 26003836',11:'Marcelo Dalbone — RA: 25001986',12:'Akiane Merbach — RA: 26027003'}
for i,text in cover.items():
 p=original[i];r=p.runs[0]if p.runs else p.add_run();r.text=text
 for rr in p.runs[1:]:rr.text=''
# Mantém capa, estilos, cabeçalhos, rodapés e configuração de página do original.
for p in original[18:]:p._p.getparent().remove(p._p)
for p in doc.paragraphs:
 for n in p._p.xpath('./w:pPr/w:numPr'):n.getparent().remove(n)

def para(text):
 p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.15
 p.paragraph_format.widow_control=True
 return p

def title(index):
 el=deepcopy(headings[index]);doc._body._body.insert(len(doc._body._body)-1,el)
 pp=el.find(qn('w:pPr'))
 if pp is None:pp=OxmlElement('w:pPr');el.insert(0,pp)
 for tag in ['numPr','pageBreakBefore','keepNext']:
  for c in pp.findall(qn('w:'+tag)):pp.remove(c)
 if index in (0,8):pp.append(OxmlElement('w:pageBreakBefore'))
 else:
  spacing=pp.find(qn('w:spacing'))
  if spacing is None:spacing=OxmlElement('w:spacing');pp.append(spacing)
  spacing.set(qn('w:before'),'280')
 pp.append(OxmlElement('w:keepNext'))

def sub(text):
 p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
 r=p.add_run(text);r.bold=True

def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'999999');borders.append(border)
 t._tbl.tblPr.append(borders)
 for c,text in zip(t.rows[0].cells,headers):
  c.text=text
  for r in c.paragraphs[0].runs:r.bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=str(text)
 for index,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1.05
    if len(t.rows)<=6 and index<len(t.rows)-1:p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.size=Pt(10)
 doc.add_paragraph()


def image(name,caption,width=5.7):
 p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(OUT/'capturas'/name),width=Inches(width))
 p=doc.add_paragraph(caption);p.paragraph_format.space_after=Pt(10)
 for r in p.runs:r.font.size=Pt(9);r.italic=True

def code(path,start,end):
 para(f'Trecho de {path}, linhas {start}–{end}:')
 text='\n'.join((ROOT/path).read_text().splitlines()[start-1:end]);p=doc.add_paragraph()
 for i,line in enumerate(text.splitlines()):
  r=p.add_run(('\n'if i else'')+line);r.font.name='Consolas';r.font.size=Pt(9)

def pair(left,right,caption):
 t=doc.add_table(rows=1,cols=2)
 for c,name in zip(t.rows[0].cells,[left,right]):
  c.paragraphs[0].add_run().add_picture(str(OUT/'capturas'/name),width=Inches(2.35))
  c.paragraphs[0].paragraph_format.keep_with_next=True
 para(caption)

title(0)
sub('Problema e público-alvo')
para('A seletividade alimentar pode limitar a variedade da alimentação infantil e dificultar a rotina dos responsáveis. Preferências por cores, texturas e preparações influenciam a aceitação dos alimentos, enquanto o acompanhamento do que foi consumido exige organização. O NutriVibe aborda esse contexto por meio de uma ferramenta de educação alimentar e registro de refeições, voltada a responsáveis e crianças.')
para('A proposta reúne informações sobre alimentos, preferências e consumo diário para apoiar a exploração de alternativas compatíveis com o perfil. A atuação é educativa: não realiza diagnóstico, tratamento ou prescrição de dieta. A presença de uma restrição no cadastro não permite garantir a ausência de ingredientes ocultos ou traços de alergênicos em uma preparação.')
sub('Solução proposta e relevância')
para('A solução é composta por uma landing page de apresentação e um aplicativo Web. A landing page explica o problema e as funcionalidades e permite cadastrar contatos mediante consentimento. O aplicativo reúne perfil, preferências, diário alimentar, resumo do dia, sugestões e conteúdo educativo. O registro pode começar pela seleção no catálogo ou por uma foto, sempre com confirmação ou correção pelo usuário.')
para('O projeto relaciona-se ao ODS 3 — Saúde e Bem-Estar — pelo acesso a informações educativas e apoio à organização da rotina. Relaciona-se também ao ODS 12 — Consumo e Produção Responsáveis — por conteúdos de planejamento e redução de desperdício. Os benefícios pretendidos são maior clareza sobre os registros e estímulo à variedade alimentar, sem prometer mudança clínica ou comportamental.')
sub('Viabilidade e valor para o usuário')
para('HTML, CSS e JavaScript permitem oferecer uma aplicação responsiva no navegador. O catálogo delimitado, a separação dos módulos e o armazenamento local viabilizam o fluxo principal sem exigir conta remota. O servidor atende serviços adicionais, como contatos. Essa organização permite desenvolver e testar as partes de maneira independente e ampliar a solução com manutenção controlada.')
para('O valor central está em reunir o que foi consumido, as preferências e as informações alimentares em um percurso compreensível. A sugestão explica suas razões e respeita os filtros cadastrados. A estimativa nutricional utiliza referência e peso confirmados, distinguindo dados conhecidos de informações ausentes.')

title(1)
sub('Identidade visual e organização das telas')
para('A identidade utiliza a marca NutriVibe em verde e amarelo, fundos sólidos, contraste e espaçamento entre controles. A landing page apresenta a solução em seções; o aplicativo organiza as ações como uma área de trabalho. No computador, a navegação é lateral; no celular, um botão abre o menu. As páginas utilizam rótulos, foco visível e mensagens de retorno.')
para('As capturas apresentam a interface em execução. O perfil Ana, de 12 anos, a restrição à lactose e a exclusão da cor verde são dados de um cenário de teste. O alimento registrado nesse cenário é Banana com referência TACO e peso de 100 g. Esses dados servem para demonstrar as interações e não correspondem a acompanhamento nutricional de uma pessoa real.')
image('landingpage.png','Figura 1 — Landing page: apresentação da proposta, navegação por seções e acesso ao aplicativo.')
image('painel.png','Figura 2 — Tela inicial do aplicativo: indicadores, registros do dia e atalhos para as funções principais.')
image('preferencias.png','Figura 3 — Perfil e preferências: formulário de identificação e opções que orientam filtros e sugestões.')
image('diario.png','Figura 4 — Diário alimentar: escolha de data e refeição e entrada do registro por foto ou catálogo.')
image('resumo.png','Figura 5 — Resumo diário: registros, cobertura e valores nutricionais disponíveis.')
image('educacao.png','Figura 6 — Educação alimentar: conteúdos de apoio e perguntas frequentes.')
sub('Fluxo de navegação e uso em celular')
para('O percurso principal é perfil e preferências, seguido de diário, seleção ou identificação do alimento, confirmação da referência e do peso e consulta ao resumo. Quando a foto não produz candidato aceitável, o catálogo e o registro manual permitem continuar. O usuário pode alterar o dia consultado, remover uma entrada e fazer backup dos dados.')
pair('painel-mobile.png','menu-mobile.png','Figura 7 — Interface em celular, com painel reorganizado e menu aberto. O botão informa seu estado e a tecla Escape permite fechar a navegação.')

title(2)
sub('Backlog e organização do trabalho')
para('O backlog organiza 32 tarefas, identificadas como TRAB-26 a TRAB-57, em três grupos: S1 — Estrutura e navegação; S2 — JavaScript e fluxo funcional; S3 — Consolidação e apresentação. As identificações e os títulos abaixo reproduzem o planejamento fornecido no Jira.')
backlog=json.loads((ROOT/'documentacoes/historico/backlog-jira.json').read_text())
for stage in backlog['sprints']:
 sub(stage['id']+' — '+stage['name'])
 para('Objetivo: '+stage['goal'])
 tasks=[t for t in backlog['tasks']if t['planned_sprint']==stage['id']]
 table(['Código','Tarefa','Prioridade'],[(t['jira_key'],t['title'],{'Highest':'Máxima','High':'Alta','Medium':'Média','Low':'Baixa'}.get(t['priority'],t['priority']))for t in tasks])
para('Os critérios comuns de aceite são comportamento demonstrável, interação funcional quando aplicável, navegação e estados de erro validados, legibilidade em celular e desktop e evidência de execução. A definição de pronto exige código integrado, ausência de erros bloqueantes, revisão e teste. As tarefas de documentação incluem capturas reais, identificação do material, PDF e preparação da demonstração oral.')
para('O registro fornecido apresenta as tarefas com status A fazer, sem responsáveis ou datas individuais definidos na ferramenta. Esses campos pertencem ao planejamento registrado e não descrevem, por si, o funcionamento do código. S1, S2 e S3 são agrupamentos por etiquetas, sem confirmação de sprints nativas. A prioridade recai sobre o fluxo de perfil, diário e resumo e suas dependências.')
sub('Scrum, Kanban e responsabilidades')
para('O Scrum organiza incrementos com objetivo, revisão e adaptação do trabalho. O Kanban utiliza A Fazer, Em Andamento e Concluído, com Revisão/Testes como etapa intermediária recomendada. Uma tarefa passa por implementação, verificação e registro de evidência antes de ser considerada pronta. Não há afirmação de que as colunas do quadro remoto tenham sido reconfiguradas.')
table(['Integrante','Responsabilidade'],[('Paulo Capreti','JavaScript, interações e integração dos fluxos.'),('Davi Pioli','HTML e acompanhamento do Kanban.'),('Victor Almeida','CSS, responsividade e identidade visual.'),('Marcelo Dalbone','Pesquisas, contextualização do problema e apoio aos requisitos.'),('Akiane Merbach','Pesquisas, referências alimentares e limitações das soluções.')])
para('O planejamento inicial concentra o escopo acadêmico em HTML, CSS e JavaScript. O backend, o banco de contatos e a análise visual são recursos adicionais da implementação. A divisão de trabalho preserva o fluxo principal como prioridade e permite revisar as contribuições em conjunto, sem substituir critérios de aceite por uma simples contagem de tarefas.')

title(3)
sub('Requisitos funcionais')
table(['ID','Requisito','Implementação e condição de uso'],[
('RF01','Cadastrar usuário e criança.','Perfil local com nome e idade; não há cadastro remoto ou dois papéis separados.'),
('RF02','Cadastrar preferências e recusas alimentares.','Restrições, cores, texturas e nomes evitados orientam os filtros.'),
('RF03','Registrar alimentos consumidos.','Entradas por data e refeição, seleção no catálogo, nome manual e remoção.'),
('RF04','Armazenar informações de alimentos e nutrientes.','Catálogo e recorte de 43 referências TACO com preparo e fonte.'),
('RF05','Analisar alimentos registrados.','Resumo de grupos e nutrientes do dia, com indicação de cobertura parcial.'),
('RF06','Recomendar complementações.','Priorização de grupos ausentes e exclusão de opções incompatíveis, sem prescrição.'),
('RF07','Disponibilizar educação alimentar.','Conteúdos de apoio e perguntas frequentes.'),
('RF08','Utilizar gamificação.','Avatar, pontos, recompensas e desafios são recursos previstos, condicionados à viabilidade.')])
para('Recursos adicionais incluem importação e exportação do diário em JSON, exclusão de dados locais, reconhecimento de foto com confirmação e cadastro consentido de contatos. O reconhecimento visual não substitui a seleção da referência e a informação do peso para calcular nutrientes.')
sub('Requisitos não funcionais')
table(['Área','Condição'],[('Usabilidade','Navegação rotulada, mensagens de retorno, foco visível e adaptação a diferentes larguras.'),('Desempenho','Resposta local para busca e filtros; comunicação do download e processamento visual.'),('Manutenibilidade','Separação entre HTML, CSS, módulos JavaScript, dados e serviços.'),('Segurança','Validação de entradas e origem, limites de requisição e consultas SQL parametrizadas.'),('Privacidade','Diário no navegador; conexão de banco e credenciais restritas ao servidor.'),('Rastreabilidade','Referências alimentares, preparo, pesos e resultados de avaliação identificados.')])
sub('Regras de negócio e restrições')
para('As recomendações consideram apenas a data selecionada, eliminam conflitos de perfil e priorizam grupos ainda não registrados. Os filtros não são relaxados para completar a lista. A estimativa nutricional exige referência do alimento escolhido e peso numérico maior que zero e não superior a 5000 g. Nutrientes ausentes continuam não informados; não são contabilizados como zero.')
para('Uma foto não confirma peso, ingredientes ocultos ou ausência de alergênicos. Preparações compostas podem ser registradas sem composição nutricional presumida. O contato da landing page exige consentimento. A proposta não diagnostica deficiências, não trata doenças e não substitui orientação de profissional de Nutrição.')

title(4)
sub('Tecnologias e estrutura')
para('O frontend utiliza HTML semântico, CSS com Grid e Flexbox e módulos JavaScript. A landing page e o aplicativo possuem estilos próprios. O backend utiliza Node.js 24; a função serverless adapta o servidor à Vercel. O banco PostgreSQL/Neon persiste contatos, enquanto o diário utiliza localStorage no navegador.')
table(['Parte','Responsabilidade'],[('html, css e javascript','Páginas, apresentação e interações, separadas entre os dois sites.'),('dados','Catálogo, referências TACO e configuração do reconhecimento.'),('servidor','Tratamento HTTP, contatos e adaptador de reconhecimento remoto.'),('api/index.mjs','Entrada serverless para o handler na Vercel.'),('ferramentas','Build, preparação de modelos e avaliação técnica.'),('testes e documentacoes','Verificações automatizadas, capturas e registros de implementação.')])
sub('Interações JavaScript e armazenamento')
para('app.js conecta eventos e elementos da interface aos módulos. engine.js normaliza buscas, verifica restrições e seleciona sugestões. store.js valida o estado e persiste profile e records na chave nutrivibe-v2. nutrition.js calcula valores proporcionais ao peso e agrega os registros por dia. Os nomes digitados são apresentados como texto por textContent, sem interpretação como marcação HTML.')
code('javascript/aplicativo/engine.js',29,31)
para('O filtro retorna somente os registros da data solicitada. Essa seleção evita que refeições de outro dia alterem o resumo consultado. As recomendações identificam grupos já consumidos, excluem conflitos e atribuem prioridade aos grupos ausentes.')
code('javascript/aplicativo/store.js',80,82)
para('Antes de salvar, validState confere estrutura e consistência. JSON.stringify converte o objeto validado em texto para o armazenamento. A importação rejeita dados inválidos; a exportação gera uma cópia para download. Não há sincronização automática do diário entre dispositivos.')
sub('Foto e confirmação nutricional')
para('A análise local utiliza SigLIP quantizado com ONNX Runtime Web. A imagem é reduzida a 224 × 224 pixels, convertida em tensor e comparada com vetores de classes. Classe negativa, similaridade baixa ou margem insuficiente levam à recusa e à correção manual. A foto permanece no aparelho nesse modo; o primeiro uso baixa aproximadamente 112 MB de modelo e runtime.')
para('A identificação fornece uma hipótese de alimento ou prato. O cálculo é separado: valor consumido = valor da referência × peso informado ÷ peso-base. Uma referência por 100 g é multiplicada por 1,5 quando o usuário informa 150 g. Dados não disponíveis na fonte permanecem ausentes e o resumo indica quantos registros contribuíram para o subtotal.')
image('confirmacao-nutricional.png','Figura 8 — Confirmação de alimento, referência TACO e peso. O usuário fornece 100 g de Banana e consulta a prévia antes de registrar.')
sub('Backend, API e banco de dados')
para('GET /api/status informa disponibilidade dos serviços. POST /api/contacts valida e grava nome, e-mail e consentimento. POST /api/identify mantém integração opcional com um provedor remoto. O modo de reconhecimento utilizado pelo aplicativo é browser-local; OpenAI e Vercel AI Gateway são alternativas dependentes de configuração e crédito, não requisitos para o processamento local.')
table(['Campo de contacts','Tipo e regra'],[('email','TEXT, chave primária; identifica o contato.'),('name','TEXT NOT NULL; nome validado pelo servidor.'),('consent_at','TEXT NOT NULL; data e hora ISO do consentimento.')])
para('Consultas parametrizadas separam os valores da instrução SQL. Se o e-mail já existe, nome e consentimento são atualizados. A conexão é obtida de POSTGRES_URL ou DATABASE_URL e aponta para o banco separado nutrivibe-production. Sem PostgreSQL, o servidor local usa SQLite. Na nuvem, não se utiliza arquivo efêmero como substituto de persistência. Essa tabela não armazena refeições ou preferências alimentares.')

title(5)
sub('Métodos e resultados')
para('A validação combina testes automatizados das regras, testes de integração, execução no navegador e comparação de reconhecimento de imagens. Foram executados 25 testes automatizados com node:test, todos aprovados. As verificações abrangem filtros, recomendações por data, importação, cálculo nutricional, políticas de reconhecimento e persistência de contatos.')
table(['Teste','Objetivo','Resultado'],[
('Restrições e seletividade','Excluir conflitos de lactose, alérgenos, cores e texturas.','Cenários aprovados; filtros não são relaxados.'),
('Recomendação por dia','Evitar influência de registros de outras datas.','Priorização e seleção diária aprovadas.'),
('Nutrição','Verificar referência compatível, peso e proporcionalidade.','Cálculos e rejeições previstos aprovados.'),
('Importação e backup','Rejeitar estruturas inválidas e preservar consistência.','Validação e fluxos de armazenamento aprovados.'),
('Backend','Conferir handler, transação parametrizada e armazenamento em nuvem.','Cenários automatizados aprovados.'),
('Navegação Chromium','Verificar formulários, diário, resumo, foto, remoção e contatos.','Percurso funcional documentado.'),
('Responsividade','Verificar as cinco telas em 360, 768 e 1280 pixels.','Sem transbordamento nos cenários verificados.')])
sub('Cenários de interface')
para('O cenário de registro utiliza Banana, referência TACO e peso de 100 g. A confirmação produz 98,2 kcal no resumo, mantém o registro após recarregar e permite abrir o diário pelo painel. A busca por maca recupera Maçã por normalização de acentos. Na configuração com lactose e cor verde excluídas, Leite e Brócolis não devem aparecer entre as sugestões compatíveis.')
image('sugestoes.png','Figura 9 — Sugestões produzidas pelo fluxo de teste, considerando os registros do dia e as restrições do perfil.')
para('As capturas de preferências, diário, confirmação e resumo correspondem ao mesmo percurso de teste. A navegação móvel foi verificada com abertura e fechamento do menu. Esses cenários demonstram comportamento técnico; não representam uma avaliação de usabilidade ou um atendimento com comunidade externa.')
sub('Reconhecimento de imagens e tratamento das falhas')
para('A comparação técnica avaliou MobileNet, CLIP e SigLIP. O modelo selecionado foi avaliado em 148 imagens. Entre 108 fotos de comida, houve 94 acertos do principal alimento ou prato, aproximadamente 87%. Na coleta adicional de 50 fotos de comida, houve 40 acertos, ou 80%. Esses valores descrevem a amostra e não equivalem à probabilidade de acerto para qualquer imagem enviada.')
para('Plantas e folhas produziram falsos positivos. A análise levou à adoção de limiares de similaridade e margem, controles negativos, recusa e confirmação manual. O tratamento de pratos compostos impede atribuir ingredientes e nutrientes presumidos. Em falha de rede ou de modelo, o catálogo e o registro manual continuam disponíveis. Uma referência inválida é rejeitada antes da persistência do cálculo.')

title(6)
sub('Acesso e execução')
para('A aplicação é hospedada na Vercel. A landing page está disponível em https://nutrivibe-gamma.vercel.app e o aplicativo em https://nutrivibe-gamma.vercel.app/app . O acesso requer navegador com JavaScript; a análise visual exige suporte a WebAssembly e conectividade para obter o modelo. O cadastro de contatos utiliza Neon.')
table(['Comando','Finalidade'],[('npm ci','Instalar dependências de acordo com package-lock.json.'),('npm run build','Preparar arquivos públicos, bundle visual e artefatos do modelo.'),('npm test','Executar a suíte automatizada.'),('npm start','Iniciar o servidor HTTP local, na porta padrão 8080.')])
para('A execução local requer Node.js 24 e npm. Os comandos são executados na raiz do projeto. A apresentação pode ser acessada em http://localhost:8080 e o aplicativo em http://localhost:8080/html/aplicativo/index.html . Os arquivos não devem ser abertos diretamente por file://, pois módulos, requisições e caminhos dos modelos dependem do servidor HTTP.')
sub('Configuração do ambiente')
para('NUTRIVIBE_VISION_PROVIDER=browser-local seleciona a análise no navegador. POSTGRES_URL ou DATABASE_URL configura o banco remoto. As variáveis são definidas no ambiente do processo ou do provedor; o código não carrega automaticamente um arquivo .env. As credenciais não são publicadas no frontend.')
para('A Vercel executa o build e serve a pasta public. A entrada api/index.mjs exporta o handler sem abrir uma porta adicional. As configurações incluem rotas, redirecionamentos, cabeçalhos e cache dos modelos. O servidor local oferece SQLite quando não há conexão PostgreSQL. O usuário dispõe de exportação para preservar o diário antes de trocar de navegador ou apagar o armazenamento local.')

title(7)
sub('Ampliação das funcionalidades')
para('As evoluções previstas incluem separar os dados do responsável e da criança, associar preferências ao perfil correspondente e ampliar a organização dos conteúdos educativos. A alteração do cadastro deve preservar registros existentes por meio de versionamento do estado, migração e testes de backup. O aprofundamento das recomendações exige regras verificáveis, justificativas compreensíveis e respeito às restrições, sem produzir prescrição clínica.')
para('A gamificação prevista na proposta utiliza avatar, pontos, recompensas e desafios de caráter educativo, condicionados à viabilidade. Esses recursos não estão implementados no aplicativo. Sua inclusão depende da estabilidade dos fluxos principais e de cuidados para evitar incentivo ao consumo forçado ou promessas de melhora clínica.')
table(['Evolução','Finalidade e cuidado necessário'],[('Cadastro responsável/criança','Representar os papéis sem perder os dados já armazenados.'),('Conteúdo educativo','Ampliar temas, linguagem acessível e identificação das fontes.'),('Catálogo e referências','Aumentar cobertura com preparações rastreáveis e dados ausentes preservados.'),('Reconhecimento','Avaliar fotos variadas, plantas, pratos compostos e aparelhos reais.'),('Desempenho','Reduzir custo de download e avaliar memória e latência em celulares.'),('Acessibilidade e uso','Avaliar teclado, contraste, tecnologias assistivas e clareza com usuários.'),('Contas e sincronização','Projetar autenticação, autorização e privacidade antes de persistir o diário remotamente.')])
sub('Manutenção e segurança')
para('Mudanças em catálogo, modelo ou limiares exigem regressão e comparação de falsos positivos e recusas. O servidor utiliza limitação em memória por processo; maior tráfego e múltiplas instâncias exigem avaliar um controle distribuído. Novas integrações devem ser comparadas por qualidade, custo e proteção dos dados.')
para('A análise de fotos continua limitada pela aparência: não confirma peso, composição completa ou alergênicos. A correção manual e a comunicação de incerteza precisam permanecer disponíveis. A avaliação com responsáveis e profissionais é necessária para verificar utilidade, adequação da linguagem e cuidados de uso. Não há resultados de impacto clínico ou comunitário mensurados.')

title(8)
para('O NutriVibe articula educação alimentar e tecnologia para apoiar a organização das refeições e a exploração de variedade com respeito às preferências. A solução reúne apresentação, perfil, diário, resumo e sugestões em um percurso que pode ser acessado pelo navegador, com adaptação para computador e celular.')
para('A implementação integra HTML, CSS, JavaScript modular, persistência local, referências TACO, reconhecimento visual confirmado e serviços de contato. Os testes verificam regras, entradas inválidas, cálculos, navegação e integração. Os resultados de reconhecimento demonstram tanto a utilidade do recurso quanto a necessidade de recusa e correção manual.')
para('O valor do projeto está na organização das informações e na transparência das orientações. A identificação de um alimento não é tratada como medida de peso ou garantia de composição, e uma sugestão de variedade não é apresentada como prescrição. Essa distinção permite oferecer apoio educativo sem atribuir ao software conclusões que dependem de avaliação profissional.')
para('O trabalho desenvolve competências de programação, estruturação de interfaces, pesquisa, modelagem de dados e gestão de tarefas. A distribuição entre JavaScript, HTML, CSS e pesquisas exige integração das contribuições, revisão e critérios de aceite. A manutenção de módulos e testes favorece aperfeiçoamentos e torna as decisões técnicas verificáveis.')
# Verifica que a capa e as imagens institucionais não foram apagadas.
assert original[1]._p.xpath('.//w:drawing|.//w:pict')
for p in doc.paragraphs:
 for n in p._p.xpath('./w:pPr/w:numPr'):n.getparent().remove(n)
 for r in p.runs:
  if r.font.size is None:r.font.size=Pt(11)
assert not doc._element.xpath('//w:numPr')
output=OUT/'NutriVibe_Preenchido.docx';doc.save(output)
with ZipFile(MODEL)as source,ZipFile(output)as result:
 for name in source.namelist():
  if name.startswith('word/media/'):
   assert source.read(name)==result.read(name),name
print(output)
